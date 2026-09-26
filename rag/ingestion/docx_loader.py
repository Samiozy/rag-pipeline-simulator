from pathlib import Path
from docx import Document as DocxDocument
from models import Document
from .base import DocumentLoader


class DOCXLoader(DocumentLoader):
    def load(self, path: Path) -> list[Document]:
        docx = DocxDocument(path)
        paragraphs = [p.text.strip() for p in docx.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        return [Document(text=text, source=path.name, metadata={"file_type": "docx"})]
