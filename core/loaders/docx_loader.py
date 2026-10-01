from pathlib import Path

from core.models import Document
from .base import DocumentLoader


class DOCXLoader(DocumentLoader):
    def load(self, path: Path) -> list[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise ImportError("DOCX support requires python-docx. Install it or upload TXT, Markdown, or PDF.") from exc
        docx = DocxDocument(path)
        paragraphs = [p.text.strip() for p in docx.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        return [Document(text=text, source=path.name, metadata={"file_type": "docx"})]
